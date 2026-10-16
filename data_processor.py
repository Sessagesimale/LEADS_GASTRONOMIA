import pandas as pd
from typing import List, Dict, Optional
import os
import re
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from utils.geo_utils import extract_province, CODE_TO_PROVINCE_NAME


class DataProcessor:
    def __init__(self):
        pass

    def _style_worksheet(self, worksheet, df: pd.DataFrame):
        """Applica stile intestazione, larghezza colonne automatica e menu a tendina su un foglio Excel."""
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for column_cells in worksheet.columns:
            length = max(len(str(cell.value or "")) for cell in column_cells)
            col_letter = column_cells[0].column_letter
            worksheet.column_dimensions[col_letter].width = min(max(length + 3, 12), 52)

        if "Esito Valutazione" in df.columns and len(df) > 0:
            val_col_idx = df.columns.get_loc("Esito Valutazione") + 1
            val_col_letter = get_column_letter(val_col_idx)
            dv = DataValidation(
                type="list",
                formula1='"Da contattare,In trattativa,Cliente acquisito,Scartato - Non in target"',
                allow_blank=True
            )
            worksheet.add_data_validation(dv)
            dv.add(f"{val_col_letter}2:{val_col_letter}{max(len(df) + 50, 100)}")

    def process_and_save(
        self,
        leads: List[Dict[str, str]],
        output_path: str,
        category_filter: str = "Tutti",
        province_filter: Optional[str] = None,
        split_by_province: bool = True,
        allowed_regions: Optional[List[str]] = None
    ) -> Optional[str]:
        if not leads:
            print("Nessun lead da salvare.")
            return None

        from utils.geo_utils import get_region_province_codes
        valid_provs_for_regions = set()
        if allowed_regions:
            for r in allowed_regions:
                valid_provs_for_regions.update(get_region_province_codes(r))

        # Assicura che ogni lead abbia la Provincia valorizzata prima del DataFrame
        clean_leads = []
        for item in leads:
            p = item.get("Provincia", "")
            r = item.get("Regione", "")
            if not p or (valid_provs_for_regions and p.upper() not in valid_provs_for_regions):
                p = extract_province(
                    str(item.get("Citta", "")),
                    str(item.get("Indirizzo", "")),
                    str(r)
                )
                if p:
                    item["Provincia"] = p

            # Se ci sono regioni specificate, scarta i lead con provincia fuori regione
            if valid_provs_for_regions and item.get("Provincia") and item.get("Provincia").upper() not in valid_provs_for_regions:
                continue

            # Pulisci eventuale "(XX)" ridondante dal campo Citta
            if item.get("Citta"):
                item["Citta"] = re.sub(r'\s*\([A-Za-z]{2}\)\s*$', '', str(item["Citta"])).strip()
            clean_leads.append(item)

        leads = clean_leads
        if not leads:
            print("Nessun lead corrispondente ai filtri geografici.")
            return None

        df = pd.DataFrame(leads)

        # 1. Filtro per escludere rigorosamente il perimetro dolciario
        sweet_terms = ["gelater", "cremer", "cioccolat", "pasticc", "dolc", "confett", "torron", "biscott"]
        pattern = "|".join(sweet_terms)
        df = df[~df["Nome"].astype(str).str.lower().str.contains(pattern, na=False, regex=True)]

        # 2. Filtro Tipologia (Ristoranti vs Negozi/Altro)
        if "Tipologia" not in df.columns:
            df["Tipologia"] = "Ristorante"

        if "ristoranti" in category_filter.lower() and "tutti" not in category_filter.lower() and "negozi" not in category_filter.lower():
            df = df[df["Tipologia"].astype(str).str.lower().str.contains("ristorant|trattori|osteri", regex=True)]
        elif "negozi" in category_filter.lower() and "tutti" not in category_filter.lower():
            df = df[~df["Tipologia"].astype(str).str.lower().str.contains("^ristorante$", regex=True)]

        # 3. Filtro Provincia opzionale
        if province_filter:
            prov_code = province_filter.strip().upper()
            df = df[df["Provincia"].astype(str).str.upper() == prov_code]

        if df.empty:
            print("Nessun lead corrispondente ai filtri impostati.")
            return None

        # 4. Rimozione duplicati basati su Nome e Citta (case-insensitive)
        df["_norm_name"] = df["Nome"].astype(str).str.lower().str.strip()
        df["_norm_city"] = df["Citta"].astype(str).str.lower().str.strip()
        df = df.drop_duplicates(subset=["_norm_name", "_norm_city"], keep="first").drop(columns=["_norm_name", "_norm_city"])

        # 5. Gestione Link: fallback reciproco tra Sito Web e Link Pagina
        if "Link Pagina" not in df.columns:
            df["Link Pagina"] = df.get("Sito Web", "")
        if "Sito Web" not in df.columns:
            df["Sito Web"] = df.get("Link Pagina", "")

        df["Sito Web"] = df["Sito Web"].fillna("").astype(str)
        df["Link Pagina"] = df["Link Pagina"].fillna("").astype(str)
        df.loc[df["Sito Web"].str.strip() == "", "Sito Web"] = df["Link Pagina"]
        df.loc[df["Link Pagina"].str.strip() == "", "Link Pagina"] = df["Sito Web"]

        if "Esito Valutazione" not in df.columns:
            df["Esito Valutazione"] = ""
        if "Note Commerciali" not in df.columns:
            df["Note Commerciali"] = ""

        # 6. Ordine desiderato delle colonne V2
        desired_columns = [
            "Nome", "Tipologia", "Regione", "Provincia", "Citta", "Indirizzo",
            "Telefono", "Email", "Sito Web", "Social", "Affinita Parmigiano",
            "Partita IVA", "Anno Bilancio", "Fatturato", "Utile/Perdita", "Addetti",
            "Link Pagina", "Fonte", "Esito Valutazione", "Note Commerciali"
        ]
        for col in desired_columns:
            if col not in df.columns:
                df[col] = ""
            else:
                df[col] = df[col].fillna("").astype(str)

        df = df[desired_columns]
        # Ordina per Regione, Provincia, Tipologia, Nome
        df = df.sort_values(by=["Regione", "Provincia", "Tipologia", "Nome"], ascending=True).reset_index(drop=True)

        # 7. Salvataggio in Excel con foglio riepilogativo + fogli divisi per Provincia
        actual_output_path = output_path
        base, ext = os.path.splitext(output_path)
        counter = 1

        while counter <= 10:
            try:
                with pd.ExcelWriter(actual_output_path, engine="openpyxl") as writer:
                    # Foglio 1: Riepilogo Completo
                    df.to_excel(writer, index=False, sheet_name="Tutti i Lead")
                    self._style_worksheet(writer.sheets["Tutti i Lead"], df)

                    # Fogli successivi: Divisione per Provincia
                    if split_by_province:
                        unique_provs = [p for p in df["Provincia"].unique() if str(p).strip()]
                        for p_code in sorted(unique_provs):
                            sub_df = df[df["Provincia"] == p_code].copy()
                            p_full = CODE_TO_PROVINCE_NAME.get(p_code, "")
                            # Nome tab max 31 caratteri per limite Excel
                            tab_name = f"{p_code} - {p_full}"[:31] if p_full else f"Prov_{p_code}"[:31]
                            tab_name = re.sub(r'[\\/*?:\[\]]', '', tab_name)
                            sub_df.to_excel(writer, index=False, sheet_name=tab_name)
                            self._style_worksheet(writer.sheets[tab_name], sub_df)

                print(f"File salvato con successo in: {actual_output_path}")
                return actual_output_path
            except PermissionError:
                actual_output_path = f"{base}_{counter}{ext}"
                counter += 1
            except Exception as e:
                print(f"Errore durante il salvataggio: {e}")
                raise e
        return actual_output_path
